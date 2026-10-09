# -*- coding: utf-8 -*-
"""scripts/privacy.py — política de minimização de dados pessoais (LGPD).

Define quais colunas e artefatos podem ser versionados no Git e fornece:

  check            varre os arquivos rastreados pelo Git e acusa violações
  public-master    gera a versão pública (minimizada) da Planilha Mestre
  minimize <arq>   remove colunas pessoais de um CSV/XLSX, guardando a cópia
                   completa em data/private/ antes de sobrescrever

Convenção de duas camadas:
  - completa (com dados pessoais)  -> local, ignorada pelo Git
  - pública (minimizada)           -> versionada, sufixo _publico

Uso:
  python scripts/privacy.py check
  python scripts/privacy.py public-master
  python scripts/privacy.py minimize "Novos Alvarás 2026/alvaras_taxi_sp_2026_consolidado.csv"
"""

from __future__ import annotations

import csv
import io
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PRIVATE_DIR = ROOT / "data" / "private"

# ---------------------------------------------------------------------------
# Política
# ---------------------------------------------------------------------------

# Colunas que NUNCA podem ser versionadas (dado pessoal identificado ou
# diretamente identificável de pessoa natural).
PII_PATTERNS: list[re.Pattern[str]] = [
    re.compile(r"(?:^|_)CPF", re.I),
    re.compile(r"Telefone_Celular|Telefone_Principal|Telefone_Secundario|Telefone_Residencial", re.I),
    re.compile(r"(?:^|_)Celular(?:$|_|\d)", re.I),
    re.compile(r"^Telefone_[12]$", re.I),
    re.compile(r"^DDD_?1$", re.I),
    re.compile(r"Telefones_Motoristas|Telefone_Motorista", re.I),
    re.compile(r"^Email$", re.I),
    re.compile(r"Email_Principal|CORREIO_ELETRONICO", re.I),
    re.compile(r"Bairro_Residencial", re.I),
    re.compile(r"Chave_Busca|Query_Busca", re.I),
    re.compile(r"Nascimento|^Idade_Minima_Estimada$", re.I),
    re.compile(r"^Genero$", re.I),
    re.compile(r"WhatsApp", re.I),
]

# Colunas pessoais adicionais, válidas apenas para arquivos específicos.
# Ex.: o CEP da Planilha Mestre é residencial; o CEP de pontos/CFCs/empresas
# é endereço institucional e permanece versionável.
EXTRA_PII_BY_FILE: dict[str, list[str]] = {
    "planilha_mestre_taxistas_sp_2026.csv": [r"^CEP$"],
    "planilha_mestre_taxistas_sp_2026.xlsx": [r"^CEP$"],
}

# Padrões proibidos em conteúdo (não apenas em cabeçalho).
CONTENT_PATTERNS: dict[str, re.Pattern[str]] = {
    "cpf_formatado": re.compile(r"\b\d{3}\.\d{3}\.\d{3}-\d{2}\b"),
    "link_whatsapp_pessoal": re.compile(r"wa\.me/5?5?\d{10,13}"),
    "link_api_whatsapp": re.compile(r"api\.whatsapp\.com/send\?phone=5?5?\d{10,11}"),
}

# Pastas fora do escopo da verificação: scripts (citam nomes de coluna por
# construção), docs (narrativa), e fontes originais imutáveis.
CHECK_SKIP_DIRS = ("scripts/", "docs/", "data/raw/", "data/external/", ".git/")
CHECK_SKIP_FILES = {"dashboard_mestre_taxistas_sp_2026.html"}
CHECK_EXTS = {".csv", ".xlsx", ".json", ".html"}


def extra_patterns(filename: str) -> list[re.Pattern[str]]:
    return [re.compile(p, re.I) for p in EXTRA_PII_BY_FILE.get(Path(filename).name, [])]


def is_pii_column(name: str, extra: list[re.Pattern[str]] | None = None) -> bool:
    """Retorna True se a coluna carrega dado pessoal não institucional."""
    name = (name or "").strip()
    if not name:
        return False
    if extra and any(p.search(name) for p in extra):
        return True
    return any(p.search(name) for p in PII_PATTERNS)


def pii_columns(columns, filename: str = "") -> list[str]:
    extra = extra_patterns(filename)
    return [c for c in columns if is_pii_column(str(c), extra)]


# ---------------------------------------------------------------------------
# Minimização
# ---------------------------------------------------------------------------


def _detect_sep(path: Path) -> str:
    with open(path, encoding="utf-8-sig", errors="ignore") as fh:
        head = fh.read(8192)
    return ";" if head.count(";") > head.count(",") else ","


def redact_text(value):
    """Remove CPF e links de WhatsApp pessoais embutidos em texto livre."""
    if not isinstance(value, str) or not value:
        return value
    for pat in CONTENT_PATTERNS.values():
        value = pat.sub("[REMOVIDO]", value)
    return value


def _redact_frame(df):
    mapper = df.map if hasattr(df, "map") else df.applymap
    return mapper(redact_text)


def minimize_csv(src: Path, dst: Path | None = None, backup: bool = True) -> list[str]:
    import pandas as pd

    dst = dst or src
    sep = _detect_sep(src)
    df = pd.read_csv(src, dtype=str, sep=sep, encoding="utf-8-sig")
    drop = [str(c) for c in df.columns if is_pii_column(str(c), extra_patterns(src.name))]
    if drop:
        df = df.drop(columns=drop)
    df = _redact_frame(df)
    if backup and dst.resolve() == src.resolve():
        _backup_private(src)
    dst.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(dst, index=False, sep=sep, encoding="utf-8-sig")
    return drop


def minimize_xlsx(src: Path, dst: Path | None = None, backup: bool = True) -> list[str]:
    import openpyxl

    dst = dst or src
    extra = extra_patterns(src.name)
    wb = openpyxl.load_workbook(src)
    dropped: list[str] = []
    for ws in wb.worksheets:
        header = [c.value for c in ws[1]]
        drop_idx = [i + 1 for i, v in enumerate(header) if is_pii_column(str(v or ""), extra)]
        if not drop_idx:
            continue
        dropped.extend(str(header[i - 1]) for i in drop_idx)
        for i in reversed(drop_idx):
            ws.delete_cols(i)
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = redact_text(cell.value)
    if backup and dst.resolve() == src.resolve():
        _backup_private(src)
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    wb.close()
    return dropped


def _backup_private(src: Path) -> Path:
    """Guarda a cópia completa (com dados pessoais) em data/private/."""
    rel = src.resolve().relative_to(ROOT)
    target = PRIVATE_DIR / rel
    if target.exists():
        return target
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(src.read_bytes())
    return target


# ---------------------------------------------------------------------------
# Verificação
# ---------------------------------------------------------------------------


def tracked_files() -> list[str]:
    out = subprocess.run(
        ["git", "-c", "core.quotepath=false", "ls-files", "-z"],
        capture_output=True,
        cwd=ROOT,
    )
    return [f.decode("utf-8") for f in out.stdout.split(b"\0") if f]


def _scan_file(rel: str) -> list[str]:
    path = ROOT / rel
    if not path.exists():
        return []
    ext = path.suffix.lower()
    problems: list[str] = []

    if ext == ".csv":
        with open(path, encoding="utf-8-sig", errors="ignore", newline="") as fh:
            head = fh.read(8192)
            delim = ";" if head.count(";") > head.count(",") else ","
            fh.seek(0)
            reader = csv.reader(fh, delimiter=delim)
            try:
                header = next(reader)
            except StopIteration:
                header = []
        cols = pii_columns(header, rel)
        if cols:
            problems.append("colunas pessoais: " + ", ".join(cols))
        text = path.read_text(encoding="utf-8-sig", errors="ignore")
        for label, pat in CONTENT_PATTERNS.items():
            n = len(pat.findall(text))
            if n:
                problems.append(f"{label}: {n} ocorrências")
    elif ext == ".xlsx":
        try:
            import zipfile

            import openpyxl

            wb = openpyxl.load_workbook(path, read_only=True)
            cols: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(max_row=1, values_only=True):
                    cols.extend(
                        c
                        for c in [str(v) for v in row if v is not None]
                        if is_pii_column(c, extra_patterns(rel))
                    )
            wb.close()
            if cols:
                problems.append("colunas pessoais: " + ", ".join(sorted(set(cols))))
            with zipfile.ZipFile(path) as zf:
                names = [
                    n
                    for n in zf.namelist()
                    if n.startswith("xl/worksheets/") or n == "xl/sharedStrings.xml"
                ]
                blob = b"".join(zf.read(n) for n in names).decode("utf-8", "ignore")
            for label, pat in CONTENT_PATTERNS.items():
                n = len(pat.findall(blob))
                if n:
                    problems.append(f"{label}: {n} ocorrências")
        except Exception as exc:  # pragma: no cover
            problems.append(f"falha ao ler xlsx: {exc}")
    elif ext in {".html", ".json"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
        for label, pat in CONTENT_PATTERNS.items():
            n = len(pat.findall(text))
            if n:
                problems.append(f"{label}: {n} ocorrências")
    return problems


def run_check() -> int:
    violations = 0
    scanned = 0
    for rel in tracked_files():
        if rel.startswith(CHECK_SKIP_DIRS) or rel in CHECK_SKIP_FILES:
            continue
        if Path(rel).suffix.lower() not in CHECK_EXTS:
            continue
        scanned += 1
        problems = _scan_file(rel)
        if problems:
            violations += 1
            print(f"VIOLAÇÃO  {rel}")
            for p in problems:
                print(f"           - {p}")
    print(f"\nverificados: {scanned} arquivos versionados | violações: {violations}")
    if violations:
        print(
            "Ação: rode `python scripts/privacy.py minimize <arquivo>` ou remova o "
            "arquivo do versionamento (git rm --cached) e o adicione ao .gitignore."
        )
    return 1 if violations else 0


# ---------------------------------------------------------------------------
# Geração da versão pública da Planilha Mestre
# ---------------------------------------------------------------------------

MASTER_FULL = ROOT / "planilha_mestre_taxistas_sp_2026.csv"
MASTER_PUBLIC = ROOT / "planilha_mestre_taxistas_sp_2026_publico.csv"


def build_public_master() -> int:
    if not MASTER_FULL.exists():
        print(f"ERRO: base completa não encontrada: {MASTER_FULL}")
        return 1
    drop = minimize_csv(MASTER_FULL, MASTER_PUBLIC, backup=False)
    size = MASTER_PUBLIC.stat().st_size
    print(f"OK  {MASTER_PUBLIC.name}  ({size / 1048576:.1f} MB)")
    print(f"    colunas removidas ({len(drop)}): {', '.join(drop)}")
    return 0


def main(argv: list[str]) -> int:
    if not argv or argv[0] not in {"check", "public-master", "minimize"}:
        print(__doc__)
        return 2
    cmd = argv[0]
    if cmd == "check":
        return run_check()
    if cmd == "public-master":
        return build_public_master()
    if cmd == "minimize":
        if len(argv) < 2:
            print("uso: python scripts/privacy.py minimize <arquivo.csv|xlsx>")
            return 2
        src = ROOT / argv[1]
        if not src.exists():
            print(f"ERRO: arquivo não encontrado: {src}")
            return 1
        drop = minimize_xlsx(src) if src.suffix.lower() == ".xlsx" else minimize_csv(src)
        print(f"OK  {src.name} — colunas removidas: {', '.join(drop) or 'nenhuma'}")
        print(f"    cópia completa preservada em data/private/{src.relative_to(ROOT)}")
        return 0
    return 2


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    raise SystemExit(main(sys.argv[1:]))
